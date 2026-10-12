"""Budget lines: the esfer@ import, charging statement movements to them, Jev's proposals."""

from __future__ import annotations

import io
import json
from datetime import date

from openpyxl import Workbook

from conftest import build_client, make_user, register_admin
from cosecre_hub.api.partides import esfera, suggest
from cosecre_hub.models import BankMovement, Document, PaymentMatch, StatementImport
from cosecre_hub.services.classification import DocumentClassifier, JevClient, JevError

API = "/api/v1/partides"


def esfera_export(exercici: int = 2026, *, drop: tuple[str, ...] = ()) -> bytes:
    """The shape of esfer@'s "Dades del pressupost" XLS, as an xlsx."""
    rows = [
        [None] * 5,
        [None, None, None, f"Data: 12/10/{exercici}", None],
        ["Dades del pressupost", None, None, f"Exercici: {exercici}", None],
        ["Nom: Escola de Prova", "Nif: S0800000A", None, "Codi: 08000000", None],
        [None, f"Any Pressupost {exercici}", "Total Ingressos 1.000,00", "Total Despeses 7.500,00", None],
        ["Tipus", "Codi partida", "Codi finalista", "Nom partida", "Import Pressupost"],
        ["Ingressos", None, None, None, None],
        [None, "I/3000001", None, "Venda de béns", 0.0],
        [None, "I/8700001", None, "Romanents d'exercicis anteriors", 1000.0],
        [None, None, None, "Total Ingressos", 1000.0],
        ["Despeses", None, None, None, None],
        [None, "D/220000100/4210", None, "Material d'oficina i didàctic no inventariable", 5000.0],
        [None, "   D/220000100/4210.01", "   MOTXILLES2526", "   25-26 Material escolar INF - MOTXILLES", 2000.0],
        [None, "   D/220000100/4210.29", "   ", "   Material oficina", 3000.0],
        [None, "D/223000100/4210", None, "Transports (autocars per a sortides curriculars)", 2500.0],
        [None, "   D/223000100/4210.11", "   FAMILIESSORTIDES2526", "   25-26 Autocars INFANTIL.", 2500.0],
        [None, "D/226003900/4210", None, "Despeses per comissions bancàries", 0.0],
        [None, None, None, "Total Despeses", 7500.0],
    ]
    book = Workbook()
    sheet = book.active
    for row in rows:
        if row[1] and row[1].strip() in drop:
            continue
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def upload(client, headers, content: bytes, *, dry_run: bool = False):
    return client.post(
        f"{API}/import",
        headers=headers,
        files={"file": ("Dades del pressupost.xlsx", content)},
        data={"dry_run": "true" if dry_run else "false"},
    )


def add_document(client, reference: str = "DOC-1", **values) -> int:
    defaults = {
        "internal_doc_number": reference,
        "tipus_document": "Factura",
        "proveidor": "Autocars Pep SL",
        "cif_proveidor": "B12345678",
        "data_factura": date(2026, 3, 10),
        "import_value": 420.0,
        "descripcio": "Servei d'autocar sortida",
        "descripcio_compra": "Autocar per la sortida d'infantil al Tibidabo",
        "status": "needs_validation",
    }
    defaults.update(values)
    with client.app.state.session_factory() as session:
        document = Document(**defaults)
        session.add(document)
        session.commit()
        return document.id


def add_movement(
    client,
    *,
    paid: str | None = None,
    concepte: str = "TRANSF. AUTOCARS PEP SL",
    import_value: float = -420.0,
    data: date = date(2026, 3, 12),
    categoria: str = "pagament",
    match: str = "confirmed",
) -> int:
    """A statement line; with ``paid``, the register entry a person confirmed it paid."""
    with client.app.state.session_factory() as session:
        statement = session.query(StatementImport).first()
        if statement is None:
            statement = StatementImport(source="caixa_xls", compte="General", file_name="x.xls")
            session.add(statement)
            session.flush()
        movement = BankMovement(
            import_id=statement.id,
            fingerprint=f"fp-{session.query(BankMovement).count() + 1}",
            source="caixa_xls",
            compte="General",
            tipus="Transferència",
            categoria=categoria,
            data=data,
            concepte=concepte,
            import_value=import_value,
            match_status="confirmed" if paid and match == "confirmed" else "unmatched",
        )
        session.add(movement)
        session.flush()
        if paid:
            document = session.query(Document).filter(Document.internal_doc_number == paid).one()
            session.add(PaymentMatch(movement_id=movement.id, document_id=document.id, status=match))
        session.commit()
        return movement.id


def lines(client, headers) -> dict[str, dict]:
    return {p["codi"]: p for p in client.get(API, headers=headers).json()["partides"]}


# ── Reading the export ──────────────────────────────────────────────────────


def test_the_parser_reads_the_tree_the_funds_and_the_header():
    budget = esfera.parse(esfera_export(), "x.xlsx")
    assert (budget.exercici, budget.centre, budget.nif, budget.codi_centre) == (
        2026,
        "Escola de Prova",
        "S0800000A",
        "08000000",
    )
    assert budget.total_despeses == 7500.0
    by_code = {p.codi: p for p in budget.partides}
    assert by_code["D/220000100/4210.01"].codi_pare == "D/220000100/4210"
    assert by_code["D/220000100/4210.01"].codi_finalista == "MOTXILLES2526"
    assert by_code["D/220000100/4210.01"].nom == "25-26 Material escolar INF - MOTXILLES"
    assert by_code["D/223000100/4210"].codi_pare is None
    assert by_code["I/3000001"].tipus == "ingres"
    assert esfera.fund_family("PlaDeXoc25_26") == ("PLADEXOC", "Pla de Xoc")
    assert esfera.fund_family("PMOEPROA25_26_FSE")[0] == "PMOEPROA"
    assert esfera.fund_family("FAM_MAT_2324_2425")[0] == "FAMILIESMATERIAL"


def test_a_file_that_is_not_the_export_is_refused_with_where_to_find_it(hub):
    client, _ = hub
    admin = register_admin(client)
    book = Workbook()
    book.active.append(["Data", "Import"])
    buffer = io.BytesIO()
    book.save(buffer)
    response = upload(client, admin, buffer.getvalue())
    assert response.status_code == 422
    assert "Consulta les dades del pressupost" in response.json()["detail"]


# ── Import ──────────────────────────────────────────────────────────────────


def test_a_dry_run_says_what_would_change_and_changes_nothing(hub):
    client, _ = hub
    admin = register_admin(client)
    preview = upload(client, admin, esfera_export(), dry_run=True).json()
    assert preview["dry_run"] is True
    assert preview["created"] == 8
    assert preview["despeses"] == 6
    assert preview["nif_matches"] is None
    assert client.get(API, headers=admin).json()["exercici"] is None

    done = upload(client, admin, esfera_export()).json()
    assert done["created"] == 8
    budget = client.get(API, headers=admin).json()
    assert budget["exercici"] == 2026
    assert budget["centre"] == "Escola de Prova"
    by_code = {p["codi"]: p for p in budget["partides"]}
    assert by_code["D/220000100/4210"]["te_fills"] is True
    assert by_code["D/220000100/4210.01"]["fons"] == "MOTXILLES"
    assert by_code["D/226003900/4210"]["te_fills"] is False


def test_reimporting_keeps_context_and_retires_only_lines_in_use(hub):
    client, _ = hub
    admin = register_admin(client)
    upload(client, admin, esfera_export())
    by_code = lines(client, admin)
    office = by_code["D/220000100/4210.29"]["id"]
    bus = by_code["D/223000100/4210.11"]["id"]
    client.put(f"{API}/lines/{office}/context", headers=admin, json={"context": "Paper i tòner"})
    movement = add_movement(client)
    assert client.put(
        f"{API}/movements/{movement}", headers=admin, json={"partides": [{"partida_id": bus}]}
    ).status_code == 200

    summary = upload(
        client, admin, esfera_export(drop=("D/223000100/4210.11", "D/226003900/4210"))
    ).json()
    assert (summary["retired"], summary["removed"], summary["unchanged"]) == (1, 1, 6)

    after = lines(client, admin)
    assert after["D/220000100/4210.29"]["context"] == "Paper i tòner"
    assert after["D/223000100/4210.11"]["retirada"] is True
    assert "D/226003900/4210" not in after


def test_a_new_year_inherits_what_was_written_about_the_same_line(hub):
    client, _ = hub
    admin = register_admin(client)
    upload(client, admin, esfera_export(2025))
    office = lines(client, admin)["D/220000100/4210.29"]["id"]
    client.put(f"{API}/lines/{office}/context", headers=admin, json={"context": "Paper i tòner"})

    summary = upload(client, admin, esfera_export(2026)).json()
    assert summary["context_inherited"] == 1
    assert lines(client, admin)["D/220000100/4210.29"]["context"] == "Paper i tòner"
    assert client.get(API, headers=admin).json()["exercicis"] == [2026, 2025]


def test_the_import_checks_the_nif_against_the_centre(hub):
    client, _ = hub
    admin = register_admin(client)
    settings = client.get("/api/v1/documents/settings", headers=admin).json()
    settings.pop("classifier_configured")
    settings.pop("drive_folder_configured")
    client.put("/api/v1/documents/settings", headers=admin, json={**settings, "school_cif": "Q0000000J"})
    assert upload(client, admin, esfera_export(), dry_run=True).json()["nif_matches"] is False


# ── Assigning ───────────────────────────────────────────────────────────────


def test_only_the_secretari_writes_the_office_reads(hub):
    client, _ = hub
    admin = register_admin(client)
    office = make_user(client, admin, "administratiu")
    upload(client, admin, esfera_export())
    movement = add_movement(client)
    line = lines(client, admin)["D/220000100/4210.29"]["id"]

    assert client.get(API, headers=office).status_code == 200
    assert client.get(f"{API}/queue", headers=office).status_code == 200
    body = {"partides": [{"partida_id": line}]}
    assert client.put(f"{API}/movements/{movement}", headers=office, json=body).status_code == 403
    assert upload(client, office, esfera_export()).status_code == 403
    assert client.put(f"{API}/context", headers=office, json={"context": "x"}).status_code == 403
    assert (
        client.put(f"{API}/lines/{line}/context", headers=office, json={"context": "x"}).status_code
        == 403
    )
    assert client.post(f"{API}/movements/{movement}/suggest", headers=office).status_code == 403


def test_a_movement_takes_several_lines_but_never_a_heading(hub):
    client, _ = hub
    admin = register_admin(client)
    upload(client, admin, esfera_export())
    movement = add_movement(client)
    by_code = lines(client, admin)

    heading = client.put(
        f"{API}/movements/{movement}",
        headers=admin,
        json={"partides": [{"partida_id": by_code["D/220000100/4210"]["id"]}]},
    )
    assert heading.status_code == 422
    assert "subpartides" in heading.json()["detail"]

    split = [
        {"partida_id": by_code["D/223000100/4210.11"]["id"], "import_assignat": 300},
        {"partida_id": by_code["D/220000100/4210.29"]["id"], "import_assignat": 120},
    ]
    item = client.put(f"{API}/movements/{movement}", headers=admin, json={"partides": split}).json()
    assert item["estat"] == "assignada"
    assert [a["import_assignat"] for a in item["assignments"]] == [300, 120]
    assert lines(client, admin)["D/223000100/4210.11"]["imputat"] == 300

    cleared = client.put(f"{API}/movements/{movement}", headers=admin, json={"partides": []}).json()
    assert cleared["estat"] == "pendent"


def test_the_queue_is_movements_with_the_invoices_they_paid(hub):
    client, _ = hub
    admin = register_admin(client)
    add_document(client, "DOC-1")
    add_document(client, "DOC-2", proveidor="Llibreria", descripcio_compra="Llibres")
    paid = add_movement(client, paid="DOC-1")
    proposed = add_movement(client, paid="DOC-2", match="proposed", data=date(2026, 3, 1))
    income = add_movement(client, concepte="QUOTES FAMILIES", import_value=150.0, categoria="ingres", data=date(2025, 10, 2))
    add_movement(client, concepte="TRASPAS A TARGETA", categoria="traspas_intern")
    add_movement(client, concepte="SALDO INICIAL", import_value=1000.0, categoria="saldo_inicial")

    queue = client.get(f"{API}/queue", headers=admin).json()
    assert [item["id"] for item in queue] == [paid, proposed, income]
    first = queue[0]
    assert (first["sentit"], first["curs"], first["concepte"]) == ("despesa", "25-26", "TRANSF. AUTOCARS PEP SL")
    assert [f["reference"] for f in first["factures"]] == ["DOC-1"]
    assert first["factures"][0]["descripcio_compra"].startswith("Autocar")
    # Only a confirmed payment says what the money was for.
    assert queue[1]["factures"] == []
    assert queue[2]["sentit"] == "ingres"


def test_deleting_a_statement_forgets_its_lines(hub):
    client, _ = hub
    admin = register_admin(client)
    upload(client, admin, esfera_export())
    movement = add_movement(client)
    line = lines(client, admin)["D/220000100/4210.29"]["id"]
    client.put(f"{API}/movements/{movement}", headers=admin, json={"partides": [{"partida_id": line}]})
    with client.app.state.session_factory() as session:
        statement_id = session.get(BankMovement, movement).import_id
    assert client.delete(f"/api/v1/statements/imports/{statement_id}", headers=admin).status_code == 204

    from cosecre_hub.api.partides.models import MovementPartida

    with client.app.state.session_factory() as session:
        assert session.query(MovementPartida).count() == 0


# ── Jev ─────────────────────────────────────────────────────────────────────


class FakeJev(JevClient):
    def __init__(self, answers=None, fail=False):
        self.model = "jev-fake"
        self.answers = answers
        self.fail = fail
        self.asked: list[tuple[dict, dict]] = []

    def ask(self, state, questions):
        if self.fail:
            raise JevError("down")
        self.asked.append((json.loads(state), questions))
        return self.answers(questions) if callable(self.answers) else dict(self.answers or {})

    def close(self):
        pass


def _jev_client(tmp_path, jev):
    client, _ = build_client(tmp_path, classifier=DocumentClassifier(jev))
    return client


def _prefer(code_ids: dict[str, int], favourite: str, *, split: float = 0.1, runner_up: str | None = None):
    def answer(questions):
        keys = list(questions["partida"]["criteria"])
        probabilities = {key: 0.0 for key in keys}
        probabilities[f"p{code_ids[favourite]}"] = 0.7
        if runner_up:
            probabilities[f"p{code_ids[runner_up]}"] = 0.25
        probabilities["cap"] = 1 - sum(probabilities.values())
        return {
            "partida": {
                "type": "choice",
                "choice": f"p{code_ids[favourite]}",
                "confidence": 0.66,
                "probabilities": probabilities,
            },
            "diverses": {"type": "noul", "noul": split},
        }

    return answer


def test_jev_proposes_a_preliminary_line_from_the_leaves_with_the_context(tmp_path):
    jev = FakeJev()
    with _jev_client(tmp_path, jev) as client:
        admin = register_admin(client)
        upload(client, admin, esfera_export())
        by_code = lines(client, admin)
        ids = {code: line["id"] for code, line in by_code.items()}
        client.put(f"{API}/context", headers=admin, json={"context": "Els autocars van a Transports."})
        client.put(
            f"{API}/lines/{ids['D/223000100/4210.11']}/context",
            headers=admin,
            json={"context": "Sortides d'infantil"},
        )
        jev.answers = _prefer(ids, "D/223000100/4210.11")
        add_document(client)
        movement = add_movement(client, paid="DOC-1")

        item = client.post(f"{API}/movements/{movement}/suggest", headers=admin).json()

    assert item["estat"] == "proposta"
    assert item["assignments"] == [
        {
            "partida_id": ids["D/223000100/4210.11"],
            "import_assignat": None,
            "origen": "jev",
            "confianca": 0.7,
            "confirmada": False,
        }
    ]
    assert item["suggestion"]["ranking"][0]["partida_id"] == ids["D/223000100/4210.11"]

    state, questions = jev.asked[0]
    assert state["curs_escolar_del_moviment"] == "25-26"
    assert state["criteris_del_centre"] == "Els autocars van a Transports."
    assert state["moviment"]["concepte"] == "TRANSF. AUTOCARS PEP SL"
    assert state["factures_que_paga"][0]["que_es_compra_segons_el_professor"].startswith("Autocar")
    criteria = questions["partida"]["criteria"]
    # Headings are not options; leaves and the way out are.
    assert f"p{ids['D/220000100/4210']}" not in criteria
    assert f"p{ids['I/3000001']}" not in criteria
    assert criteria[f"p{ids['D/223000100/4210.11']}"]["quan_s_usa"] == "Sortides d'infantil"
    assert criteria[f"p{ids['D/223000100/4210.11']}"]["fons"] == "Famílies · sortides"
    assert "cap" in criteria
    assert questions["diverses"]["type"] == "noul"


def test_jev_never_overwrites_what_a_person_confirmed(tmp_path):
    jev = FakeJev()
    with _jev_client(tmp_path, jev) as client:
        admin = register_admin(client)
        upload(client, admin, esfera_export())
        ids = {code: line["id"] for code, line in lines(client, admin).items()}
        jev.answers = _prefer(ids, "D/223000100/4210.11")
        movement = add_movement(client)
        client.put(
            f"{API}/movements/{movement}",
            headers=admin,
            json={"partides": [{"partida_id": ids["D/220000100/4210.29"]}]},
        )
        item = client.post(f"{API}/movements/{movement}/suggest", headers=admin).json()

    assert [a["partida_id"] for a in item["assignments"]] == [ids["D/220000100/4210.29"]]
    assert item["assignments"][0]["confirmada"] is True
    assert item["suggestion"] is not None


def test_a_split_purchase_gets_two_lines_and_the_supplier_history_is_a_hint(tmp_path):
    jev = FakeJev()
    with _jev_client(tmp_path, jev) as client:
        admin = register_admin(client)
        upload(client, admin, esfera_export())
        ids = {code: line["id"] for code, line in lines(client, admin).items()}
        add_document(client, "DOC-OLD", data_factura=date(2026, 1, 5))
        old = add_movement(client, paid="DOC-OLD", data=date(2026, 1, 9))
        client.put(
            f"{API}/movements/{old}",
            headers=admin,
            json={"partides": [{"partida_id": ids["D/223000100/4210.11"]}]},
        )
        jev.answers = _prefer(ids, "D/223000100/4210.11", split=0.8, runner_up="D/220000100/4210.29")
        add_document(client, "DOC-1", cif_proveidor="B-12345678")
        movement = add_movement(client, paid="DOC-1")
        item = client.post(f"{API}/movements/{movement}/suggest", headers=admin).json()

    assert [a["partida_id"] for a in item["assignments"]] == [
        ids["D/223000100/4210.11"],
        ids["D/220000100/4210.29"],
    ]
    state, _ = jev.asked[0]
    assert state["partides_anteriors_del_mateix_proveidor_o_concepte"] == [
        {"partida": "25-26 Autocars INFANTIL.", "vegades": 1}
    ]


def test_without_jev_or_a_budget_the_suggestion_says_why(tmp_path, hub):
    client, _ = hub
    admin = register_admin(client)
    movement = add_movement(client)
    assert client.post(f"{API}/movements/{movement}/suggest", headers=admin).status_code == 503

    with _jev_client(tmp_path / "jev", FakeJev(fail=True)) as other:
        (tmp_path / "jev").mkdir(exist_ok=True)
        admin = register_admin(other)
        movement = add_movement(other)
        assert other.post(f"{API}/movements/{movement}/suggest", headers=admin).status_code == 409
        upload(other, admin, esfera_export())
        assert other.post(f"{API}/movements/{movement}/suggest", headers=admin).status_code == 502


def test_a_school_year_runs_september_to_august():
    assert suggest.school_year(date(2025, 9, 1)) == "25-26"
    assert suggest.school_year(date(2026, 8, 31)) == "25-26"
    assert suggest.school_year(None) == ""


def test_a_split_adds_up_to_the_movement_or_is_refused(hub):
    client, _ = hub
    admin = register_admin(client)
    upload(client, admin, esfera_export())
    movement = add_movement(client)  # −420 €
    by_code = lines(client, admin)
    bus, office = by_code["D/223000100/4210.11"]["id"], by_code["D/220000100/4210.29"]["id"]

    def put(*parts):
        body = {"partides": [{"partida_id": pid, "import_assignat": amount} for pid, amount in parts]}
        return client.put(f"{API}/movements/{movement}", headers=admin, json=body)

    short = put((bus, 300), (office, 100))
    assert short.status_code == 422
    assert "420.00" in short.json()["detail"]
    assert put((bus, 300), (office, None)).status_code == 422
    assert put((bus, 420), (office, 0)).status_code == 422
    assert put((bus, 315.5), (office, 104.5)).status_code == 200
    # Even, unsplit: allowed, and counted half each.
    assert put((bus, None), (office, None)).status_code == 200
    assert lines(client, admin)["D/223000100/4210.11"]["imputat"] == 210
    # One line takes the whole movement whatever it was sent with.
    alone = put((bus, 12)).json()
    assert alone["assignments"][0]["import_assignat"] is None


def test_money_in_is_offered_the_income_lines_and_bank_fees_their_concept(tmp_path):
    jev = FakeJev()
    with _jev_client(tmp_path, jev) as client:
        admin = register_admin(client)
        upload(client, admin, esfera_export())
        ids = {code: line["id"] for code, line in lines(client, admin).items()}
        jev.answers = _prefer(ids, "I/3000001")
        income = add_movement(client, concepte="INGRES QUOTES", import_value=80.0, categoria="ingres")
        client.post(f"{API}/movements/{income}/suggest", headers=admin)

        fee = add_movement(client, concepte="COMISSIO MANTENIMENT", import_value=-6.0, categoria="comissio")
        client.put(
            f"{API}/movements/{fee}", headers=admin, json={"partides": [{"partida_id": ids["D/226003900/4210"]}]}
        )
        jev.answers = _prefer(ids, "D/226003900/4210")
        again = add_movement(client, concepte="Comissió manteniment", import_value=-6.0, categoria="comissio")
        client.post(f"{API}/movements/{again}/suggest", headers=admin)

    state, questions = jev.asked[0]
    criteria = questions["partida"]["criteria"]
    assert state["moviment"]["sentit"] == "entrada de diners"
    assert f"p{ids['I/3000001']}" in criteria
    assert f"p{ids['D/220000100/4210.29']}" not in criteria
    state, _ = jev.asked[1]
    assert state["partides_anteriors_del_mateix_proveidor_o_concepte"] == [
        {"partida": "Despeses per comissions bancàries", "vegades": 1}
    ]


def test_the_secretari_orders_a_concepts_lines_and_it_holds(hub):
    client, _ = hub
    admin = register_admin(client)
    office = make_user(client, admin, "administratiu")
    upload(client, admin, esfera_export())
    body = {"block": "D/220000100/4210", "codis": ["D/220000100/4210.29", "D/220000100/4210.01"]}
    assert client.put(f"{API}/order", headers=office, json=body).status_code == 403
    assert client.put(f"{API}/order", headers=admin, json=body).json() == {
        "D/220000100/4210": ["D/220000100/4210.29", "D/220000100/4210.01"]
    }
    assert client.get(API, headers=office).json()["ordre_usuari"]["D/220000100/4210"][0] == "D/220000100/4210.29"
    cleared = client.put(f"{API}/order", headers=admin, json={"block": "D/220000100/4210", "codis": []})
    assert cleared.json() == {}


def test_jev_only_preselects_when_it_is_sure():
    unsure = suggest.Suggestion(
        ranking=[suggest.Ranked(1, 0.45), suggest.Ranked(2, 0.3)], confidence=0.45, cap=0.1, multiple=0.1, model="m"
    )
    assert unsure.preassign() == []
    sure = suggest.Suggestion(ranking=[suggest.Ranked(1, 0.8)], confidence=0.8, cap=0.1, multiple=0.1, model="m")
    assert [r.partida_id for r in sure.preassign()] == [1]
