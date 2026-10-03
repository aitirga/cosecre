"""CaixaBank's "Moviments del compte" export, read without a model.

The bank's export never changes shape, so this is plain parsing:

* a title row, ``Moviments del compte ES00 2100 … (CCC: …)``, that names the
  account — which is how a file tells General, Material i Sortides and
  Menjador apart;
* a header row ``Data | Data valor | Moviment | Més dades | Import | Saldo``,
  found by its labels rather than its position;
* one row per movement, newest first, with Excel date serials and signed
  amounts.

What a row *is* (transfer, direct debit, card, fee…) is not in the file; it is
read off the concept the way a person would, by :func:`classify`.

Run it on a file to see what it reads::

    uv run python -m cosecre_hub.services.statements.caixa_xls statement.xls
"""

from __future__ import annotations

import io
import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path
from typing import Any

from ..text_format import iban_is_valid, normalize_iban, parse_amount, parse_date
from . import (
    FEE,
    INCOME,
    INTERNAL,
    PAYMENT,
    SOURCE_CAIXA,
    ParsedMovement,
    ParsedStatement,
    StatementError,
)

TRANSFER = "Transferència bancària"
DIRECT_DEBIT = "Rebut domiciliat"
DEBIT_CARD = "Targeta de dèbit"

_HEADERS = {
    "data": "data",
    "data valor": "data_valor",
    "moviment": "moviment",
    "mes dades": "mes_dades",
    "import": "import",
    "saldo": "saldo",
}
_IBAN = re.compile(r"\b([A-Z]{2}\d{2}(?:\s?[A-Z0-9]{4}){4,7}(?:\s?[A-Z0-9]{1,4})?)")
#: ``FAC:PRF26-01609``, ``Fac: ES 139/2026``. The bank cuts the concept at about
#: seventeen characters, so what follows can be the start of a longer number.
_INVOICE = re.compile(r"^\s*FAC(?:TURA)?\s*[:.]\s*(.+?)\s*$", re.IGNORECASE)
#: A SEPA creditor reference: the creditor's CIF/NIF plus a three-digit suffix.
_CREDITOR = re.compile(r"^([A-HJ-NP-SUVW]\d{8}|\d{8}[A-Z])(\d{3})$")
_FEES = ("MANTENIMENT", "CORRESP.", "COMISSIO", "COMISION", "INTERESSOS", "LIQUIDACIO", "DESPESES")
_TOP_UP = ("CARREGA.TARG.PREPAG", "CARREGA TARG", "RECARGA TARJETA")
_TRANSFER_WORDS = ("TRANSF", "TRASPAS")


def _fold(text: str) -> str:
    decomposed = unicodedata.normalize("NFKD", str(text).lower())
    return " ".join("".join(c for c in decomposed if not unicodedata.combining(c)).split())


def classify(moviment: str, mes_dades: str, amount: float) -> dict[str, str]:
    """What a row is, from what the bank printed. Unknown payments are card ones.

    Returns ``tipus``, ``categoria`` and whatever hints the text carries
    (``num_factura_hint``, ``cif_hint``).
    """
    concept = (moviment or "").strip().upper()
    extra = (mes_dades or "").strip()
    extra_compact = re.sub(r"\s", "", extra).upper()
    result = {"tipus": DEBIT_CARD, "categoria": PAYMENT, "num_factura_hint": "", "cif_hint": ""}

    if any(concept.startswith(word) for word in _TOP_UP):
        return {**result, "tipus": TRANSFER, "categoria": INTERNAL}
    if amount > 0:
        return {**result, "tipus": TRANSFER, "categoria": INCOME}
    if any(concept.startswith(word) for word in _FEES):
        return {**result, "tipus": "", "categoria": FEE}
    if match := _INVOICE.match(concept):
        return {**result, "tipus": TRANSFER, "num_factura_hint": match.group(1)}
    if match := _CREDITOR.match(extra_compact):
        return {**result, "tipus": DIRECT_DEBIT, "cif_hint": match.group(1)}
    if "REBUT" in extra_compact or "RECIBO" in extra_compact or concept.startswith(("REBUT", "RECIBO")):
        return {**result, "tipus": DIRECT_DEBIT}
    if any(word in concept for word in _TRANSFER_WORDS):
        return {**result, "tipus": TRANSFER}
    # A card purchase leaves "Més dades" empty; a transfer with a free-text
    # concept ("Reserva P2026076") names who received it there.
    if re.search(r"[A-Za-z]{3}", extra):
        return {**result, "tipus": TRANSFER}
    return result


def read_rows(content: bytes, file_name: str = "") -> list[list[Any]]:
    """Every cell of the first sheet, as plain Python values."""
    is_xlsx = content[:2] == b"PK" or file_name.lower().endswith((".xlsx", ".xlsm"))
    try:
        if is_xlsx:
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = book.worksheets[0]
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        import xlrd

        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_index(0)
        return [sheet.row_values(r) for r in range(sheet.nrows)]
    except Exception as exc:  # noqa: BLE001 — any reader failure means "not this format"
        raise StatementError("No s'ha pogut llegir el full de càlcul.") from exc


def _date(value: Any) -> date | None:
    # xlrd gives serials as floats; openpyxl gives datetimes; both go through here.
    return parse_date(value)


def parse(content: bytes, file_name: str = "") -> ParsedStatement:
    rows = read_rows(content, file_name)

    iban = ""
    header_index = None
    columns: dict[str, int] = {}
    for index, row in enumerate(rows[:15]):
        cells = [str(c or "").strip() for c in row]
        joined = " ".join(cells)
        if not iban and (match := _IBAN.search(joined.upper())):
            candidate = normalize_iban(match.group(1))
            if iban_is_valid(candidate):
                iban = candidate
        folded = [_fold(c) for c in cells]
        if "moviment" in folded and "import" in folded:
            header_index = index
            columns = {_HEADERS[f]: i for i, f in enumerate(folded) if f in _HEADERS}
            break
    if header_index is None or "data" not in columns:
        raise StatementError(
            "Aquest Excel no té el format de La Caixa (Data, Moviment, Més dades, Import, Saldo)."
        )

    def cell(row: list[Any], name: str) -> Any:
        i = columns.get(name)
        return row[i] if i is not None and i < len(row) else None

    movements: list[ParsedMovement] = []
    for row in rows[header_index + 1 :]:
        when = _date(cell(row, "data"))
        amount = parse_amount(cell(row, "import"))
        if when is None or amount is None:
            continue  # a blank or a footer line
        moviment = str(cell(row, "moviment") or "").strip()
        mes_dades = str(cell(row, "mes_dades") or "").strip()
        saldo = parse_amount(cell(row, "saldo"))
        value_date = _date(cell(row, "data_valor"))
        kind = classify(moviment, mes_dades, amount)
        movements.append(
            ParsedMovement(
                data=when,
                data_valor=value_date,
                concepte=moviment,
                mes_dades=mes_dades,
                import_value=round(amount, 2),
                saldo=round(saldo, 2) if saldo is not None else None,
                identity=(SOURCE_CAIXA, iban, when, value_date, moviment, mes_dades, round(amount, 2), saldo),
                raw={
                    "data": when.isoformat(),
                    "data_valor": value_date.isoformat() if value_date else None,
                    "moviment": moviment,
                    "mes_dades": mes_dades,
                    "import": amount,
                    "saldo": saldo,
                },
                **kind,
            )
        )
    if not movements:
        raise StatementError("L'extracte no té cap moviment.")
    return ParsedStatement(source=SOURCE_CAIXA, movements=movements, account_iban=iban)


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print("usage: python -m cosecre_hub.services.statements.caixa_xls <file.xls>", file=sys.stderr)
        return 2
    path = Path(argv[1])
    statement = parse(path.read_bytes(), path.name)
    print(
        json.dumps(
            {
                "account_iban": statement.account_iban,
                "period": [d.isoformat() if d else None for d in statement.period],
                "movements": [
                    {
                        "data": m.data.isoformat() if m.data else None,
                        "concepte": m.concepte,
                        "mes_dades": m.mes_dades,
                        "import": m.import_value,
                        "tipus": m.tipus,
                        "categoria": m.categoria,
                        "num_factura_hint": m.num_factura_hint,
                        "cif_hint": m.cif_hint,
                    }
                    for m in statement.movements
                ],
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
